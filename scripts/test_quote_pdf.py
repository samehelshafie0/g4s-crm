"""End-to-end renderer checks, including actual LibreOffice DOCX/XLSX conversion."""
import importlib.util
import io
import subprocess
import shutil
import tempfile
import unittest
from pathlib import Path

from docx import Document
from openpyxl import Workbook
from PIL import Image
from pypdf import PdfReader, PdfWriter
from reportlab.pdfgen.canvas import Canvas

ROOT = Path(__file__).resolve().parents[1]
spec = importlib.util.spec_from_file_location('renderer', ROOT/'crm-api/internal/services/pdf_renderer.py')
renderer = importlib.util.module_from_spec(spec)
spec.loader.exec_module(renderer)


def fixture(directory, kind, content='Specification page'):
    path = directory / ('source.'+kind)
    if kind == 'pdf':
        canvas = Canvas(str(path)); canvas.drawString(45,760,content+' one'); canvas.showPage(); canvas.drawString(45,760,content+' two'); canvas.save()
        reader = PdfReader(path); reader.pages[1].rotate(90)
        writer = PdfWriter(); writer.append(reader)
        with open(path,'wb') as f: writer.write(f)
    elif kind == 'docx':
        doc = Document(); doc.add_heading('Customer Agreement',0); doc.add_paragraph('Agreement first page. Customer-specific delivery conditions.'); doc.add_page_break(); doc.add_heading('Acceptance',1); doc.add_paragraph('Agreement second page.'); doc.save(path)
    elif kind == 'xlsx':
        book = Workbook(); first = book.active; first.title = 'Equipment'; first.append(['Specification', 'Quantity']); first.append(['Camera',2]); second = book.create_sheet('Maintenance'); second.append(['Service','Visits']); second.append(['Inspection',12])
        for sheet in book:
            sheet.column_dimensions['A'].width=35; sheet.column_dimensions['B'].width=15
            sheet.print_area='A1:B8'; sheet.page_setup.paperSize=sheet.PAPERSIZE_A4
            sheet.page_setup.orientation='landscape'; sheet.page_setup.fitToWidth=1; sheet.sheet_properties.pageSetUpPr.fitToPage=True
        book.save(path)
    elif kind == 'png': Image.new('RGB',(1000,600),(30,100,170)).save(path)
    return path


def quote(rows=65):
    return {'quoteNumber':'QT-PDF-ACCEPTANCE','status':'approved','currency':'SAR','customer':{'companyName':'PDF Acceptance Customer'},'soldTo':{'company':'PDF Acceptance Customer','contactName':'Procurement Team'},'shipTo':{},'validUntil':'2027-01-01','discountPercent':0,'discountAmount':0,'vatPercent':15,'vatAmount':rows*15,'subtotal':rows*100,'subtotalAfterDiscount':rows*100,'total':rows*115,'introductionText':'Please find our quotation and the supporting specifications and customer agreement.','paymentTerms':'Net 30','deliveryTerms':'As agreed','statementOfWork':'Installation, testing and commissioning.','notes':'SENSITIVE_GENERIC_CRM_NOTE','internalNotes':'SENSITIVE_INTERNAL_PRIVATE','purchasingNotes':'SENSITIVE_PURCHASING_PRIVATE','totalCost':987654.321,'lineItems':[{'rowType':'item','description':f'Security equipment item {i+1}. Includes installation and commissioning.','sku':f'SKU-{i+1:03}','quantity':2,'multiplier':1,'unitCost':987654.321,'unitPrice':50,'discountPercent':0,'lineTotal':100,'isPrintable':True,'isSelected':True} for i in range(rows)]}


class PDFTests(unittest.TestCase):
    def test_multi_page_quote_with_pdf_word_excel_and_image(self):
        with tempfile.TemporaryDirectory() as folder:
            directory = Path(folder)
            items = []
            for kind,label in [('pdf','Technical specifications'),('docx','Customer agreement'),('xlsx','Equipment and maintenance schedule'),('png','Site reference image')]:
                path=fixture(directory,kind)
                items.append({'path':str(path),'fileName':path.name,'label':label,'version':'2'})
            output = directory/'combined.pdf'
            renderer.render({'quote':quote(),'appendices':items},output)
            reader=PdfReader(output); texts=[page.extract_text() for page in reader.pages]
            combined='\n'.join(texts)
            base_count=len(reader.pages)-sum(item['pages'] for item in items)
            self.assertGreaterEqual(base_count,4)
            self.assertEqual([item['pages'] for item in items],[2,2,2,1])
            self.assertIn('Grand total',combined); self.assertIn('7,475.00',combined)
            self.assertNotIn('SENSITIVE_',combined); self.assertNotIn('987,654',combined)
            start=base_count
            for i,item in enumerate(items,1):
                for n in range(item['pages']):
                    self.assertIn(f'Appendix {i}:',texts[start+n]); self.assertIn(f'Document page {n+1} of {item["pages"]}',texts[start+n])
                start+=item['pages']
            self.assertIn('Specification page one',texts[base_count])
            self.assertIn('Specification page two',texts[base_count+1])
            self.assertIn('Agreement first page',texts[base_count+2])
            self.assertIn('Equipment',combined)
            for n,text in enumerate(texts,1): self.assertIn(f'Page {n} of {len(reader.pages)}',text)
            self.assertEqual(len(reader.outline),5)
            target=ROOT/'tmp/pdfs/qa-quote-with-appendices.pdf'; target.parent.mkdir(parents=True,exist_ok=True); target.write_bytes(output.read_bytes())
            print(f'QA artifact: {target} ({base_count} quote pages + 7 appendix pages)')

    def test_cropped_content_stays_hidden_and_form_fields_survive(self):
        import pymupdf
        from pypdf.generic import RectangleObject
        with tempfile.TemporaryDirectory() as folder:
            directory=Path(folder); source=directory/'cropped.pdf'
            canvas=Canvas(str(source),pagesize=(400,400))
            canvas.setFillColorRGB(1,0,0); canvas.rect(0,360,400,40,fill=1,stroke=0); canvas.rect(0,0,400,30,fill=1,stroke=0)
            canvas.setFillColorRGB(0,0,0);canvas.drawString(50,200,'Visible specification')
            canvas.acroForm.textfield(name='customer',value='Visible customer',x=50,y=250,width=200,height=25)
            canvas.acroForm.textfield(name='cropped',value='Hidden widget',x=50,y=365,width=200,height=25)
            canvas.showPage();canvas.save()
            writer=PdfWriter();writer.append(PdfReader(source));writer.pages[0].cropbox=RectangleObject((0,40,400,350))
            with source.open('wb') as out:writer.write(out)
            result=directory/'result.pdf'
            renderer.render({'quote':quote(1),'appendices':[{'path':str(source),'fileName':'cropped.pdf','version':'1','label':'Cropped specification'}]},result)
            reader=PdfReader(result);fields=reader.get_fields()
            self.assertEqual(fields['appendix_1.customer']['/V'],'Visible customer')
            self.assertEqual(fields['appendix_1.cropped']['/V'],'Hidden widget')
            rendered=pymupdf.open(result);pix=rendered[-1].get_pixmap()
            image=Image.frombytes('RGB',(pix.width,pix.height),pix.samples)
            red_pixels=sum(1 for r,g,b in image.get_flattened_data() if r>200 and g<60 and b<60)
            self.assertEqual(red_pixels,0,'Cropped red content became visible in appendix margins')
            self.assertTrue(any(widget.field_value=='Visible customer' for widget in rendered[-1].widgets()))

    def test_legacy_word_and_excel_formats(self):
        with tempfile.TemporaryDirectory() as folder:
            directory=Path(folder)
            office=shutil.which('soffice') or '/Applications/LibreOffice.app/Contents/MacOS/soffice'
            items=[]
            for modern,legacy in [('docx','doc'),('xlsx','xls')]:
                source=fixture(directory,modern)
                profile=directory/('legacy-profile-'+legacy)
                subprocess.run([office,'-env:UserInstallation='+profile.as_uri(),'--headless','--convert-to',legacy,'--outdir',str(directory),str(source)],check=True,stdout=subprocess.DEVNULL,stderr=subprocess.DEVNULL,timeout=60)
                old=source.with_suffix('.'+legacy)
                self.assertTrue(old.exists())
                items.append({'path':str(old),'fileName':old.name,'label':'Legacy '+legacy.upper(),'version':'1'})
            output=directory/'legacy.pdf'; renderer.render({'quote':quote(1),'appendices':items},output)
            self.assertEqual([item['pages'] for item in items],[2,2])
            combined=''.join(p.extract_text() for p in PdfReader(output).pages)
            self.assertIn('Agreement first page',combined); self.assertIn('Legacy XLS',combined)

    def test_no_appendices_and_excluded_optional_content(self):
        with tempfile.TemporaryDirectory() as folder:
            q=quote(1); q['lineItems'].append({**q['lineItems'][0],'description':'EXCLUDED_OPTION','isOptional':True,'isSelected':False})
            out=Path(folder)/'quote.pdf'; renderer.render({'quote':q,'appendices':[]},out)
            text=''.join(page.extract_text() for page in PdfReader(out).pages)
            self.assertNotIn('EXCLUDED_OPTION',text); self.assertNotIn('Included documents',text)

    def test_encrypted_and_invalid_files_fail_without_partial_export(self):
        with tempfile.TemporaryDirectory() as folder:
            directory=Path(folder); path=fixture(directory,'pdf'); writer=PdfWriter(); writer.append(PdfReader(path)); writer.encrypt('secret')
            locked=directory/'locked.pdf'
            with locked.open('wb') as f: writer.write(f)
            item={'path':str(locked),'fileName':'locked.pdf','label':'Locked specification','version':'1'}
            out=directory/'out.pdf'
            with self.assertRaisesRegex(ValueError,'Password-protected'): renderer.render({'quote':quote(1),'appendices':[item]},out)
            self.assertFalse(out.exists())
            locked.write_text('not a PDF')
            with self.assertRaisesRegex(ValueError,'Appendix 1'): renderer.render({'quote':quote(1),'appendices':[item]},out)
            self.assertFalse(out.exists())


if __name__=='__main__': unittest.main()
